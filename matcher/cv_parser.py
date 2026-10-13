"""Read a CV and pull out contact details, skills, and target roles."""

from __future__ import annotations

import re
import zipfile
from io import BytesIO
from pathlib import Path

SKILLS = [
    "Python", "Java", "JavaScript", "TypeScript", "C#", "C++", "Rust", "PHP", "Ruby",
    "Swift", "Kotlin", "SQL", "NoSQL", "HTML", "CSS", "React", "Angular", "Vue",
    "Node.js", "Next.js", "Django", "Flask", "FastAPI", "Fastify", "Spring", ".NET", "ASP.NET", "Laravel",
    "Docker", "Kubernetes", "AWS", "Azure", "GCP", "Linux", "Git", "REST", "GraphQL",
    "MongoDB", "PostgreSQL", "MySQL", "Redis", "Prisma", "Tailwind CSS", "Bootstrap",
    "Material UI", "Ant Design", "jQuery", "Zod", "Postman", "Tableau", "Power BI", "Excel",
    "Looker", "Figma", "Adobe Photoshop", "Adobe Illustrator", "Jira", "Salesforce",
    "SAP", "Oracle", "TensorFlow", "PyTorch", "Machine Learning", "Deep Learning",
    "NLP", "Computer Vision", "Pandas", "NumPy", "scikit-learn", "Statistics",
    "A/B testing", "ETL", "Airflow", "Snowflake", "BigQuery", "Terraform", "CI/CD",
    "Agile", "Scrum", "Kafka", "Elasticsearch", "PowerPoint", "Microsoft Office",
    "Google Sheets", "AutoCAD", "SolidWorks", "MATLAB", "SPSS", "Project Management",
    "Product Management", "Business Analysis", "Financial Analysis", "Accounting",
    "Bookkeeping", "IFRS", "Auditing", "Budgeting", "Forecasting", "Procurement",
    "Supply Chain", "Logistics", "Customer Service", "Human Resources", "Recruiting",
    "Talent Acquisition", "Payroll", "SEO", "SEM", "Google Analytics", "Copywriting",
    "Social Media", "CRM", "Negotiation", "Market Research", "Business Development",
    "Stakeholder Management", "QuickBooks", "Risk Management", "Compliance",
    "UX Research", "Wireframing", "Prototyping", "Nursing", "Patient Care",
    "Clinical Research", "Teaching", "Curriculum Development", "Translation",
    "Customer Success", "Technical Support", "Quality Assurance", "Manual Testing",
    "Test Automation", "Cypress", "Selenium", "Unit Testing", "Data Analysis",
    "Data Science", "Data Engineering", "Business Intelligence", "Financial Modeling",
    "Microsoft Excel", "Customer Support", "Account Management", "Digital Marketing",
    "Content Marketing", "Email Marketing", "Sales Management", "Cold Calling",
    "Public Speaking", "Grant Writing", "Legal Research", "Contract Review",
    "Inventory Management", "Warehouse Management", "Electrical Engineering",
    "Mechanical Engineering", "Civil Engineering", "Project Planning",
]

ALIASES = {
    "js": "JavaScript",
    "ts": "TypeScript",
    "nodejs": "Node.js",
    "node": "Node.js",
    "reactjs": "React",
    "react.js": "React",
    "react js": "React",
    "nextjs": "Next.js",
    "next js": "Next.js",
    "html5": "HTML",
    "css3": "CSS",
    "rest api": "REST",
    "rest apis": "REST",
    "mui": "Material UI",
    "tailwind": "Tailwind CSS",
    "tailwindcss": "Tailwind CSS",
    "postgres": "PostgreSQL",
    "postgresql": "PostgreSQL",
    "mongodb": "MongoDB",
    "powerbi": "Power BI",
    "power-bi": "Power BI",
    "ml": "Machine Learning",
    "ai": "Machine Learning",
    "ab testing": "A/B testing",
    "a/b tests": "A/B testing",
    "ms excel": "Excel",
    "microsoft excel": "Excel",
    "google bigquery": "BigQuery",
    "nodejs.js": "Node.js",
}

AMBIGUOUS = {
    "Sales", "Marketing", "Design", "Support", "Management", "Research", "Training",
    "Service", "Finance", "Writing", "Planning", "Operations", "Strategy", "Content",
    "Human Resources", "Translation", "Budgeting", "Quality Assurance",
}

SECTION_ONLY = {"Go", "R", "C"}

HEADINGS = {
    "summary": "summary",
    "professional summary": "summary",
    "profile": "summary",
    "about me": "summary",
    "about": "summary",
    "objective": "summary",
    "skills": "skills",
    "technical skills": "skills",
    "core competencies": "skills",
    "competencies": "skills",
    "technologies": "skills",
    "tools": "skills",
    "expertise": "skills",
    "languages": "skills",
    "experience": "experience",
    "work experience": "experience",
    "work experiences": "experience",
    "professional experience": "experience",
    "professional experiences": "experience",
    "employment": "experience",
    "work history": "experience",
    "education": "other",
    "projects": "other",
    "certifications": "other",
    "contact": "other",
}

TITLE_HINT = re.compile(
    r"\b(engineer|engineering|developer|designer|manager|analyst|scientist|"
    r"consultant|specialist|coordinator|assistant|director|intern|officer|"
    r"architect|administrator|accountant|nurse|teacher|recruiter|marketer|"
    r"marketing|writer|editor|lawyer|attorney|salesperson|representative|"
    r"technician|mechanic|supervisor|associate|executive|founder|researcher|"
    r"pharmacist|therapist|chef|planner|buyer|programmer|president|"
    r"administrator|head)\b",
    re.I,
)

ACTION_START = re.compile(
    r"^(led|built|managed|improved|developed|created|worked|responsible|helped|"
    r"achieved|increased|reduced|designed|implemented|supported|delivered|"
    r"collaborated|assisted|organized|prepared|maintained|used|using|"
    r"responsible)\b",
    re.I,
)

EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)
PHONE_RE = re.compile(r"(?:\+\d{1,3}[\s.-]?)?(?:\(?\d{2,4}\)?[\s.-]?)?\d{3,4}[\s.-]\d{3,4}(?:[\s.-]\d{2,4})?")
LINKEDIN_RE = re.compile(r"https?://(?:[\w.]+)?linkedin\.com/in/[\w\-_%]+/?", re.I)
LOCATION_RE = re.compile(r"^(?:location|address|city)\s*[:|-]\s*(.+)$", re.I)
PLACE_RE = re.compile(r"\b([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*,\s*[A-Z][a-z]+)\b")
SKILL_LABEL_RE = re.compile(
    r"^(?:languages|language|technologies|technology|tools|frameworks|libraries|databases|database|other|soft skills)\s*:\s*",
    re.I,
)
SKILL_STOP = {"apis", "api", "beginner", "present", "part-time", "part time", "skills", "technologies"}


def _canonical_map() -> dict[str, str]:
    mapping = {skill.lower(): skill for skill in SKILLS}
    mapping.update(ALIASES)
    for skill in SECTION_ONLY:
        mapping[skill.lower()] = skill
    return mapping


CANONICAL = _canonical_map()

_TERMS = sorted(CANONICAL, key=len, reverse=True)
_SKILL_RE = re.compile(
    r"(?<!\w)(?:" + "|".join(re.escape(term) for term in _TERMS) + r")(?!\w)",
    re.I,
)


def read_upload(filename: str, data: bytes) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ValueError(
                    "This PDF is password-protected. Export an unlocked copy or paste the text."
                ) from exc
        if len(reader.pages) > 30:
            raise ValueError("That PDF has too many pages for a CV. Paste the text instead.")
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    elif ext == ".docx":
        assert_docx_limits(data)
        from docx import Document

        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
        text = "\n".join(parts)
    elif ext == ".txt":
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError("Upload a PDF, DOCX, or TXT file.")

    text = normalize_text(text)
    if len(text) < 40:
        raise ValueError(
            "Could not read enough text from that file. If it is a scanned PDF, paste the CV text instead."
        )
    return text


def assert_docx_limits(data: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            total = 0
            for info in archive.infolist():
                total += info.file_size
                if info.file_size > 20 * 1024 * 1024 or total > 20 * 1024 * 1024:
                    raise ValueError("That DOCX is too large to read. Paste the CV text instead.")
    except zipfile.BadZipFile as exc:
        raise ValueError("That file could not be read. Try PDF, DOCX, or paste the text.") from exc


def normalize_text(text: str) -> str:
    text = text.replace("\ufb01", "fi").replace("\ufb02", "fl")
    text = text.replace("\u00a0", " ").replace("\u2013", "-").replace("\u2014", "-")
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_upload(filename: str, data: bytes) -> dict:
    text = read_upload(filename, data)
    profile = parse_cv(text)
    profile["cv_filename"] = Path(filename).name
    return profile


def parse_cv(text: str) -> dict:
    text = normalize_text(text)
    lines = [clean_line(line) for line in text.splitlines()]
    lines = [line for line in lines if line]
    sections = split_sections(lines)

    skills = extract_skills(text, sections.get("skills", []))
    titles = extract_titles(lines, sections.get("experience", []))
    summary = extract_summary(sections.get("summary", []), lines)
    email = first_match(EMAIL_RE, text)
    linkedin = first_match(LINKEDIN_RE, text)
    phone = extract_phone(text)
    location = extract_location(lines)
    name = extract_name(lines, email)
    if name.isupper():
        name = name.title()

    return {
        "name": name,
        "email": email,
        "phone": phone,
        "location": location,
        "linkedin": linkedin,
        "summary": summary[:600],
        "skills": skills[:40],
        "target_titles": titles[:6],
        "cv_text": text[:200000],
        "cv_filename": "",
    }


def clean_line(line: str) -> str:
    line = line.strip()
    line = re.sub(r"^[\s•\-\*\u2022\u00b7]+", "", line)
    return re.sub(r"\s+", " ", line).strip()


def split_sections(lines: list[str]) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"preamble": []}
    current = "preamble"
    for line in lines:
        heading = heading_name(line)
        if heading:
            current = heading
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return sections


def heading_name(line: str) -> str | None:
    if len(line) > 40:
        return None
    cleaned = re.sub(r"[:\-\s]+$", "", line).strip().lower()
    return HEADINGS.get(cleaned)


def extract_skills(text: str, skill_lines: list[str]) -> list[str]:
    found: list[str] = []
    seen: set[str] = set()

    def add(skill: str) -> None:
        skill = " ".join(skill.split())
        skill = re.sub(r"\s*\([^)]*\)", "", skill).strip(" .-:")
        if not skill or skill.lower() in SKILL_STOP:
            return
        key = skill.lower()
        canon = CANONICAL.get(key, skill)
        if canon.lower() in seen:
            return
        seen.add(canon.lower())
        found.append(canon)

    section_text = "\n".join(skill_lines)
    beginner = beginner_skills(section_text)
    for token in split_skill_tokens(section_text):
        if token.lower() in beginner:
            continue
        add(token)

    counts: dict[str, int] = {}
    for match in _SKILL_RE.finditer(text):
        canon = CANONICAL[match.group(0).lower()]
        if canon in SECTION_ONLY:
            continue
        counts[canon] = counts.get(canon, 0) + 1

    section_keys = {item.lower() for item in found}
    for canon, count in counts.items():
        if canon.lower() in beginner:
            continue
        if canon in AMBIGUOUS and count < 2 and canon.lower() not in section_keys:
            continue
        add(canon)
    return found


def beginner_skills(text: str) -> set[str]:
    found = set()
    for item in re.split(r"[,;/|\n•·]", text):
        if not re.search(r"\bbeginner\b", item, re.I):
            continue
        token = SKILL_LABEL_RE.sub("", " ".join(item.split()))
        token = re.sub(r"\s*\([^)]*\)", "", token).strip(" .-:")
        if not token:
            continue
        found.add(CANONICAL.get(token.lower(), token).lower())
    return found


def split_skill_tokens(text: str) -> list[str]:
    if not text:
        return []
    raw = re.split(r"[,;/|\n•·]|(?:\s+\band\b\s+)", text, flags=re.I)
    tokens = []
    for item in raw:
        token = SKILL_LABEL_RE.sub("", " ".join(item.split())).strip(" .-:")
        token = re.sub(r"\s*\([^)]*\)", "", token).strip(" .-:")
        if not token or len(token) < 2 or len(token) > 40:
            continue
        if len(token.split()) > 5:
            continue
        if "@" in token or "http" in token.lower():
            continue
        if token.lower() in SKILL_STOP or token.lower() in {"and", "or", "the", "with", "for", "from"}:
            continue
        if re.search(r"\b(?:19|20)\d{2}\b", token) or heading_name(token):
            continue
        tokens.append(token)
    return tokens


def extract_titles(lines: list[str], experience_lines: list[str]) -> list[str]:
    titles: list[str] = []
    seen: set[str] = set()

    def add(candidate: str) -> None:
        candidate = tidy_title(candidate)
        if not title_like(candidate):
            return
        key = candidate.lower()
        if key in seen:
            return
        seen.add(key)
        titles.append(candidate)

    for line in lines[:8]:
        consider_title_line(line, add)
        if titles:
            break

    for line in experience_lines:
        consider_title_line(line, add)
    return titles


def tidy_title(candidate: str) -> str:
    candidate = re.sub(r"\s*\([^)]*\)", "", candidate)
    candidate = re.split(r"[,/|]", candidate)[0]
    return " ".join(candidate.split()).strip(" -|")


def consider_title_line(line: str, add) -> None:
    if "|" in line or " - " in line:
        parts = re.split(r"\s+[|\-]\s+", line)
        matched = [part for part in parts if title_like(part)]
        if matched:
            add(matched[0])
            return
    add(line)


def title_like(line: str) -> bool:
    line = " ".join(line.split()).strip()
    if not line or len(line) < 4 or len(line) > 80:
        return False
    words = line.split()
    if not 1 < len(words) <= 6:
        return False
    if "@" in line or "%" in line or line.count(",") > 1:
        return False
    if re.search(r"\d", line):
        return False
    if ACTION_START.match(line):
        return False
    if heading_name(line):
        return False
    return TITLE_HINT.search(line) is not None


def extract_summary(summary_lines: list[str], lines: list[str]) -> str:
    if summary_lines:
        text = " ".join(summary_lines)
        return text[:600]
    for line in lines:
        if len(line) >= 80 and not title_like(line) and "@" not in line:
            return line[:600]
    return ""


def extract_name(lines: list[str], email: str) -> str:
    for line in lines[:6]:
        if email and email in line:
            continue
        if "@" in line or "http" in line.lower() or ":" in line:
            continue
        words = line.split()
        if not 1 < len(words) <= 4:
            continue
        if len(line) > 40:
            continue
        if all(re.fullmatch(r"[A-Za-z][A-Za-z.'-]*", word) for word in words):
            return line
    return ""


def extract_phone(text: str) -> str:
    match = PHONE_RE.search(text)
    if not match:
        return ""
    phone = match.group(0).strip()
    digits = re.sub(r"\D", "", phone)
    if len(digits) < 8 or len(digits) > 15:
        return ""
    return phone


def extract_location(lines: list[str]) -> str:
    for line in lines[:20]:
        match = LOCATION_RE.match(line)
        if match:
            return match.group(1).strip()[:80]
    for line in lines[:12]:
        if "@" not in line and "linkedin" not in line.lower():
            continue
        found = PLACE_RE.search(line)
        if found:
            return found.group(1).strip()[:80]
    return ""


def first_match(pattern: re.Pattern, text: str) -> str:
    match = pattern.search(text)
    return match.group(0).strip() if match else ""
